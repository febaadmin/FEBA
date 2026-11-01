"""
Fabrique les gabarits bureautiques de TEST (versionnés dans ce dossier).

Ce sont de vrais fichiers Word, OpenDocument, RTF et PowerPoint, avec ce
qu'un gabarit officiel contient réellement : orientation, bordure de page
colorée, bandeau de couleur, logo (image), tableau de signatures, texte
accentué, et pour certains des marqueurs {{ champ }}.

    pip install python-docx==1.1.2 python-pptx==1.0.2   # outils de fabrication seulement
    python tests/fixtures/office_templates/fabriquer.py
    # puis, dans un conteneur qui a LibreOffice (formats dérivés : DOC, ODT, RTF, PPT, ODP) :
    docker run --rm -v "$PWD/tests/fixtures/office_templates:/f" -w /f feba-backend \\
        sh -c 'for t in odt doc rtf; do soffice --headless --convert-to $t diplome_paysage_marqueurs.docx; done; \\
               for t in ppt odp; do soffice --headless --convert-to $t diplome_diapo_marqueurs.pptx; done'

Les fichiers DOC, ODT, RTF, PPT et ODP sont donc produits par LibreOffice
à partir des DOCX / PPTX : ce sont des fichiers de ces formats, que FEBA
reçoit tels quels et doit reconvertir lui-même. Aucun PDF n'est fourni aux
tests : la conversion est toujours faite par FEBA.
"""
import io
import os

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor as PptColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Mm as PMm, Pt as PPt

HERE = os.path.dirname(os.path.abspath(__file__))


def logo_png(color=(176, 138, 46)):
    image = Image.new("RGB", (400, 400), (255, 255, 255))
    draw = ImageDraw.Draw(image)
    draw.ellipse([20, 20, 380, 380], fill=color, outline=(30, 58, 110), width=18)
    draw.rectangle([150, 120, 250, 280], fill=(30, 58, 110))
    out = io.BytesIO()
    image.save(out, "PNG")
    out.seek(0)
    return out


def page_border(section, color="B08A2E", size=36):
    """Bordure de page (w:pgBorders), comme un cadre de diplôme."""
    sect_pr = section._sectPr
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "page")
    for side in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:val"), "double")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "24")
        element.set(qn("w:color"), color)
        borders.append(element)
    sect_pr.append(borders)


def shade(cell, hex_color):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), hex_color)
    properties.append(shading)


def paragraph(document, text, size=14, bold=False, color=(30, 58, 110), align="center",
              italic=False, space_after=6):
    p = document.add_paragraph()
    p.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(*color)
    return p


def diploma_docx(path, markers=True):
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Mm(297), Mm(210)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Mm(22))
    section.top_margin = section.bottom_margin = Mm(18)
    page_border(section)

    head = document.add_table(rows=1, cols=1)
    head.alignment = WD_TABLE_ALIGNMENT.CENTER
    shade(head.rows[0].cells[0], "1E3A6E")
    run = head.rows[0].cells[0].paragraphs[0].add_run("ACADÉMIE D'ESSAI - RÉPUBLIQUE DES TESTS")
    head.rows[0].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    run.font.color.rgb = RGBColor(255, 255, 255)
    run.font.bold = True
    run.font.size = Pt(12)

    logo = document.add_paragraph()
    logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo.add_run().add_picture(logo_png(), width=Mm(22))
    paragraph(document, "DIPLÔME DE FIN D'ÉTUDES", size=30, bold=True, space_after=2)
    paragraph(document, "décerné à", size=14, italic=True, color=(90, 90, 90))
    paragraph(document, "{{ student_full_name }}" if markers else " ", size=24, bold=True,
              color=(176, 138, 46))
    paragraph(document, ("pour l'année scolaire {{ school_year }} - {{ academy_name }}"
                         if markers else "pour l'année scolaire en cours"), size=13)

    table = document.add_table(rows=2, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells,
                          ("Le Directeur", "Sceau de l'établissement", "Délivré le")):
        cell.paragraphs[0].add_run(text).font.bold = True
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade(cell, "F3E9D2")
    for cell, text in zip(table.rows[1].cells,
                          ("", "", "{{ issue_date }}" if markers else "")):
        cell.paragraphs[0].add_run(text)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph(document, "N° {{ document_number }}" if markers else "N°", size=9,
              color=(120, 120, 120), align="left")
    document.save(path)


def certificate_docx(path):
    """Certificat PORTRAIT, sans marqueur : champs placés dans l'éditeur."""
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(20)
    page_border(section, color="1F6E46", size=24)
    logo = document.add_paragraph()
    logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo.add_run().add_picture(logo_png((31, 110, 70)), width=Mm(25))
    paragraph(document, "CERTIFICAT DE SCOLARITÉ", size=26, bold=True, color=(31, 110, 70))
    paragraph(document, "Le chef d'établissement soussigné certifie que l'élève :", size=13,
              align="left", space_after=60)
    paragraph(document, "est régulièrement inscrit(e) en classe de :", size=13, align="left",
              space_after=40)
    table = document.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    for row, label in zip(table.rows, ("Année scolaire", "Matricule", "Référence")):
        row.cells[0].paragraphs[0].add_run(label).font.bold = True
        shade(row.cells[0], "E3F1E8")
    paragraph(document, "Fait pour servir et valoir ce que de droit.", size=12, italic=True)
    document.save(path)


def multipage_docx(path):
    """Deux pages : la première est une notice, la SECONDE porte les marqueurs."""
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    paragraph(document, "NOTICE - page 1, à ne pas imprimer", size=20, bold=True)
    paragraph(document, "Cette page explique l'usage du modèle.", size=12)
    document.add_page_break()
    paragraph(document, "ATTESTATION - page 2", size=24, bold=True, color=(120, 30, 30))
    paragraph(document, "Élève : {{ student_name }}", size=16, align="left")
    paragraph(document, "Classe : {{ class_name }}", size=14, align="left")
    document.save(path)


def diploma_pptx(path):
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = PMm(297), PMm(210)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PMm(8), PMm(8), PMm(281), PMm(194))
    frame.fill.background()
    frame.line.color.rgb = PptColor(0x6A, 0x1B, 0x9A)
    frame.line.width = PPt(8)
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PMm(8), PMm(8), PMm(281), PMm(25))
    band.fill.solid()
    band.fill.fore_color.rgb = PptColor(0x6A, 0x1B, 0x9A)
    band.line.fill.background()
    slide.shapes.add_picture(logo_png((106, 27, 154)), PMm(135), PMm(38), PMm(27), PMm(27))

    def text(content, top, size, color=PptColor(0x20, 0x20, 0x20), bold=False):
        box = slide.shapes.add_textbox(PMm(30), PMm(top), PMm(237), PMm(16))
        paragraph_ = box.text_frame.paragraphs[0]
        paragraph_.alignment = 2  # centré
        run = paragraph_.add_run()
        run.text = content
        run.font.size = PPt(size)
        run.font.bold = bold
        run.font.color.rgb = color

    text("DIPLÔME D'HONNEUR", 70, 32, PptColor(0x6A, 0x1B, 0x9A), True)
    text("{{ student_name }}", 100, 26)
    text("{{ academy_name }} - {{ school_year }}", 125, 16)
    text("Délivré le {{ issue_date }}", 170, 12)
    presentation.save(path)


if __name__ == "__main__":
    diploma_docx(os.path.join(HERE, "diplome_paysage_marqueurs.docx"))
    certificate_docx(os.path.join(HERE, "certificat_portrait.docx"))
    multipage_docx(os.path.join(HERE, "deux_pages_marqueurs_page2.docx"))
    diploma_pptx(os.path.join(HERE, "diplome_diapo_marqueurs.pptx"))
    print("fichiers écrits dans", HERE)
